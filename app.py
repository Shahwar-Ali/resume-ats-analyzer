import json
import os
import re
from typing import Optional, Tuple
import docx
import pypdf
import streamlit as st
from google import genai
from google.genai import types

st.set_page_config(
    page_title="AI Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)


def extract_text_from_pdf(file_obj) -> str:
    try:
        reader = pypdf.PdfReader(file_obj)
        text_blocks = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(text_blocks).strip()
    except Exception as e:
        st.error(f"Error reading PDF: {e}")
        return ""


def extract_text_from_docx(file_obj) -> str:
    try:
        doc = docx.Document(file_obj)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs).strip()
    except Exception as e:
        st.error(f"Error reading DOCX: {e}")
        return ""


def sanitize_latex(text: str) -> str:
    if not isinstance(text, str):
        return text
    return text.replace("$", "\\$")


def analyze_resume_with_gemini(
    resume_text: str, api_key: str, model_name: str
) -> Tuple[Optional[dict], Optional[str]]:
    client = genai.Client(api_key=api_key)

    system_instruction = (
        "You are an executive ATS (Applicant Tracking System) auditor and resume specialist. "
        "Analyze the provided resume rigorously and return a strict JSON response. "
        "Each score must be an integer between 0 and 100. "
        "Do not invent contact info; report issues factually."
    )

    prompt = f"""
Analyze this resume for ATS optimization, clarity, structure, and professional impact.

Resume Content:
\"\"\"
{resume_text}
\"\"\"

Return ONLY valid JSON matching this exact structure:
{{
  "scores": {{
    "formatting_structure": 0-100,
    "action_verbs_quantification": 0-100,
    "skills_presentation": 0-100,
    "section_completeness": 0-100,
    "clarity_conciseness": 0-100
  }},
  "strengths": [
    "strength 1",
    "strength 2",
    "strength 3"
  ],
  "weaknesses": [
    "weakness 1",
    "weakness 2",
    "weakness 3"
  ],
  "actionable_recommendations": [
    "clear improvement 1",
    "clear improvement 2",
    "clear improvement 3"
  ],
  "missing_or_weak_keywords": [
    "keyword 1",
    "keyword 2"
  ]
}}
"""

    try:
        response = client.models.generate_content(
            model=model_name,
            contents=prompt,
            config=types.GenerateContentConfig(
                system_instruction=system_instruction,
                response_mime_type="application/json",
                temperature=0.2,
            ),
        )

        response_text = response.text.strip()
        cleaned_json = re.sub(
            r"^```(?:json)?\s*|\s*```$", "", response_text, flags=re.MULTILINE
        ).strip()
        parsed_data = json.loads(cleaned_json)
        return parsed_data, None

    except Exception as err:
        return None, str(err)


def calculate_weighted_overall_score(scores: dict) -> int:
    weights = {
        "formatting_structure": 0.25,
        "action_verbs_quantification": 0.25,
        "skills_presentation": 0.20,
        "section_completeness": 0.15,
        "clarity_conciseness": 0.15,
    }

    total = 0.0
    for key, weight in weights.items():
        score = scores.get(key, 70)
        score = max(0, min(100, int(score)))
        total += score * weight

    return int(round(total))


# Sidebar
st.sidebar.title("⚙️ Configuration")

default_key = ""
if "GEMINI_API_KEY" in st.secrets:
    default_key = st.secrets["GEMINI_API_KEY"]

api_key = st.sidebar.text_input(
    "Google Gemini API Key",
    value=default_key,
    type="password",
    help="Get a key at https://aistudio.google.com/apikey",
)

model_choice = st.sidebar.selectbox(
    "Gemini Model",
    options=["gemini-3.8-flash", "gemini-3.5-flash", "gemini-3.1-flash-lite"],
    index=0,
)

st.sidebar.markdown("---")
st.sidebar.caption(
    "💡 *Your uploaded files are processed in-memory and are never stored on disk.*"
)

# Main UI
st.title("📄 AI Resume ATS Analyzer")
st.write(
    "Upload your resume in PDF or DOCX format to receive an ATS score, keyword audit, and tactical improvement steps."
)

uploaded_file = st.file_uploader(
    "Upload Resume (Max 5MB)",
    type=["pdf", "docx"],
    help="Select a PDF or DOCX file to analyze.",
)

if uploaded_file is not None:
    if uploaded_file.size > 5 * 1024 * 1024:
        st.error("File size exceeds 5MB limit. Please upload a smaller document.")
        st.stop()

    file_extension = uploaded_file.name.split(".")[-1].lower()

    with st.spinner("Extracting resume contents..."):
        if file_extension == "pdf":
            resume_text = extract_text_from_pdf(uploaded_file)
        elif file_extension == "docx":
            resume_text = extract_text_from_docx(uploaded_file)
        else:
            resume_text = ""

    if not resume_text:
        st.warning(
            "⚠️ Could not extract readable text from this file. If this is a scanned PDF, please upload a text-based document."
        )
        st.stop()

    with st.expander("Preview Extracted Text"):
        st.text_area(
            "Raw Content",
            value=resume_text[:2000] + ("..." if len(resume_text) > 2000 else ""),
            height=150,
            disabled=True,
        )

    if st.button("🚀 Analyze Resume", type="primary"):
        if not api_key:
            st.error(
                "Please provide a Gemini API key in the sidebar or via secrets.toml."
            )
            st.stop()

        with st.spinner("Auditing resume against ATS standards with Gemini..."):
            result, error_msg = analyze_resume_with_gemini(
                resume_text, api_key, model_choice
            )

        if error_msg:
            st.error(f"Analysis failed: {error_msg}")
            st.info("Tip: Verify that your API key is valid and has sufficient quota.")
        elif result:
            scores = result.get("scores", {})
            overall_score = calculate_weighted_overall_score(scores)

            st.markdown("---")
            st.subheader("📊 ATS Evaluation Overview")

            col_score, col_metrics = st.columns([1, 2])

            with col_score:
                st.metric(
                    label="Overall ATS Score",
                    value=f"{overall_score} / 100",
                    delta=(
                        "Ready to Apply"
                        if overall_score >= 80
                        else (
                            "Needs Work"
                            if overall_score >= 60
                            else "High Risk of Rejection"
                        )
                    ),
                )
                if overall_score >= 80:
                    st.success("Strong resume! Passes standard ATS thresholds.")
                elif overall_score >= 60:
                    st.warning("Moderate score. Address the key issues below.")
                else:
                    st.error(
                        "Low score. Critical ATS formatting or content gaps detected."
                    )

            with col_metrics:
                m1, m2 = st.columns(2)
                m1.progress(
                    scores.get("formatting_structure", 0) / 100,
                    text=f"Formatting: {scores.get('formatting_structure', 0)}%",
                )
                m1.progress(
                    scores.get("skills_presentation", 0) / 100,
                    text=f"Skills Section: {scores.get('skills_presentation', 0)}%",
                )
                m2.progress(
                    scores.get("action_verbs_quantification", 0) / 100,
                    text=f"Impact & Metrics: {scores.get('action_verbs_quantification', 0)}%",
                )
                m2.progress(
                    scores.get("clarity_conciseness", 0) / 100,
                    text=f"Clarity: {scores.get('clarity_conciseness', 0)}%",
                )

            st.markdown("---")

            col_str, col_weak = st.columns(2)

            with col_str:
                st.subheader("✅ Key Strengths")
                for s in result.get("strengths", []):
                    st.markdown(f"- {sanitize_latex(s)}")

            with col_weak:
                st.subheader("⚠️ ATS Red Flags & Weaknesses")
                for w in result.get("weaknesses", []):
                    st.markdown(f"- {sanitize_latex(w)}")

            st.markdown("---")

            col_rec, col_kw = st.columns([2, 1])

            with col_rec:
                st.subheader("🎯 Step-by-Step Improvements")
                for i, rec in enumerate(
                    result.get("actionable_recommendations", []), 1
                ):
                    st.markdown(f"**{i}.** {sanitize_latex(rec)}")

            with col_kw:
                st.subheader("🔍 Missing / Recommended Keywords")
                keywords = result.get("missing_or_weak_keywords", [])
                if keywords:
                    for kw in keywords:
                        st.markdown(f"`{sanitize_latex(kw)}`")
                else:
                    st.write("No major keyword gaps identified.")
